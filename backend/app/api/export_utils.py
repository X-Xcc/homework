from io import BytesIO
from pydantic import BaseModel


class ExportRequest(BaseModel):
    document_name: str
    summary: str | None = ""
    risks: list[dict] = []
    overall_risk_level: str | None = None
    overall_score: float | None = None
    format: str = "txt"


def _risk_label(level: str) -> str:
    labels = {"high": "高风险", "medium": "中风险", "low": "低风险"}
    return labels.get(level, level)


def _build_report_content(data: ExportRequest) -> str:
    lines = [
        "合同风险分析报告",
        "=" * 40,
        f"文档名称: {data.document_name}",
    ]
    if data.overall_risk_level:
        lines.append(f"风险等级: {_risk_label(data.overall_risk_level)}")
    if data.overall_score is not None:
        lines.append(f"综合评分: {data.overall_score}")
    lines.extend(["", "--- 分析总结 ---", data.summary or "无", "", f"--- 风险条款 ({len(data.risks)} 项) ---"])
    for i, risk in enumerate(data.risks):
        lines.extend([
            "", f"{i+1}. [{_risk_label(risk.get('level', ''))}] {risk.get('title', '')}",
            f"   描述: {risk.get('description', '')}",
        ])
        if risk.get("location"):
            lines.append(f"   位置: {risk['location']}")
        if risk.get("legal_basis"):
            lines.append(f"   法律依据: {', '.join(risk['legal_basis'])}")
        if risk.get("suggestion"):
            lines.append(f"   建议: {risk['suggestion']}")
    lines.extend(["", f"报告生成时间: {__import__('datetime').datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
    return "\n".join(lines)


def _generate_txt(data: ExportRequest) -> BytesIO:
    content = _build_report_content(data)
    buf = BytesIO()
    buf.write(content.encode("utf-8"))
    buf.seek(0)
    return buf


def _generate_docx(data: ExportRequest) -> BytesIO:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    import datetime as _dt

    doc = Document()
    style = doc.styles["Normal"]
    font = style.font
    font.name = "SimSun"
    font.size = Pt(10.5)

    title = doc.add_heading("合同风险分析报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"文档名称: {data.document_name}")
    if data.overall_risk_level:
        doc.add_paragraph(f"风险等级: {_risk_label(data.overall_risk_level)}")
    if data.overall_score is not None:
        doc.add_paragraph(f"综合评分: {data.overall_score}")
    doc.add_paragraph("")

    doc.add_heading("分析总结", level=1)
    doc.add_paragraph(data.summary or "无")

    doc.add_heading(f"风险条款 ({len(data.risks)} 项)", level=1)
    for i, risk in enumerate(data.risks):
        doc.add_heading(f"{i+1}. [{_risk_label(risk.get('level', ''))}] {risk.get('title', '')}", level=2)
        doc.add_paragraph(f"描述: {risk.get('description', '')}")
        if risk.get("location"):
            doc.add_paragraph(f"位置: {risk['location']}")
        if risk.get("legal_basis"):
            doc.add_paragraph(f"法律依据: {', '.join(risk['legal_basis'])}")
        if risk.get("suggestion"):
            p = doc.add_paragraph()
            run = p.add_run(f"建议: {risk['suggestion']}")
            run.bold = True

    doc.add_paragraph("")
    doc.add_paragraph(f"报告生成时间: {_dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _generate_pdf(data: ExportRequest) -> BytesIO:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import os as _os
    import datetime as _dt

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=20*mm, bottomMargin=20*mm)
    styles = getSampleStyleSheet()

    _cn_font = "Helvetica"
    _cn_font_bold = "Helvetica-Bold"
    for _font_path, _font_name in [
        ("C:/Windows/Fonts/simsun.ttc", "SimSun"),
        ("C:/Windows/Fonts/simhei.ttf", "SimHei"),
        ("C:/Windows/Fonts/msyh.ttc", "Microsoft YaHei"),
    ]:
        if _os.path.exists(_font_path):
            try:
                pdfmetrics.registerFont(TTFont(_font_name, _font_path))
                _cn_font = _font_name
                _cn_font_bold = _font_name
                break
            except Exception:
                continue

    title_style = ParagraphStyle("CNTitle", parent=styles["Title"], fontName=_cn_font_bold, fontSize=18, alignment=TA_CENTER, spaceAfter=12)
    heading_style = ParagraphStyle("CNHeading", parent=styles["Heading2"], fontName=_cn_font_bold, fontSize=14, spaceBefore=12, spaceAfter=6)
    subheading_style = ParagraphStyle("CNSubHeading", parent=styles["Heading3"], fontName=_cn_font_bold, fontSize=11, spaceBefore=8, spaceAfter=4)
    body_style = ParagraphStyle("CNBody", parent=styles["Normal"], fontName=_cn_font, fontSize=10, leading=16, spaceAfter=4)

    story = []
    story.append(Paragraph("合同风险分析报告", title_style))
    story.append(Spacer(1, 6*mm))

    story.append(Paragraph(f"文档名称: {data.document_name}", body_style))
    if data.overall_risk_level:
        story.append(Paragraph(f"风险等级: {_risk_label(data.overall_risk_level)}", body_style))
    if data.overall_score is not None:
        story.append(Paragraph(f"综合评分: {data.overall_score}", body_style))
    story.append(Spacer(1, 4*mm))

    story.append(HRFlowable(width="100%", thickness=0.5))
    story.append(Paragraph("分析总结", heading_style))
    story.append(Paragraph((data.summary or "无").replace("\n", "<br/>"), body_style))

    story.append(HRFlowable(width="100%", thickness=0.5))
    story.append(Paragraph(f"风险条款 ({len(data.risks)} 项)", heading_style))

    for i, risk in enumerate(data.risks):
        level_text = _risk_label(risk.get("level", ""))
        story.append(Paragraph(f"{i+1}. [{level_text}] {risk.get('title', '')}", subheading_style))
        story.append(Paragraph(f"描述: {risk.get('description', '')}", body_style))
        if risk.get("location"):
            story.append(Paragraph(f"位置: {risk['location']}", body_style))
        if risk.get("legal_basis"):
            story.append(Paragraph(f"法律依据: {', '.join(risk['legal_basis'])}", body_style))
        if risk.get("suggestion"):
            story.append(Paragraph(f"<b>建议: {risk['suggestion']}</b>", body_style))
        story.append(Spacer(1, 2*mm))

    story.append(HRFlowable(width="100%", thickness=0.5))
    story.append(Paragraph(f"报告生成时间: {_dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", body_style))

    doc.build(story)
    buf.seek(0)
    return buf
